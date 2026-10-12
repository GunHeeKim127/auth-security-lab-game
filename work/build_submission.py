from pathlib import Path
import json
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
ASSETS = ROOT / "docs" / "assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

summary = json.loads((ROOT / "results" / "repeated-experiment-summary.json").read_text(encoding="utf-8"))
source = (ROOT / "docs" / "submission-paper-source.md").read_text(encoding="utf-8")

font_path = Path("C:/Windows/Fonts/malgun.ttf")
font_name = "Malgun Gothic"

chart_path = ASSETS / "attack-success-rate.png"
modes = summary["modes"]
labels = ["취약", "인증·인가 부분 강화", "전체 강화"]
values = [item["attackSuccessRatePercent"] for item in modes]
image = Image.new("RGB", (1200, 680), "white")
draw = ImageDraw.Draw(image)
title_font = ImageFont.truetype(str(font_path), 38) if font_path.exists() else ImageFont.load_default()
label_font = ImageFont.truetype(str(font_path), 30) if font_path.exists() else ImageFont.load_default()
value_font = ImageFont.truetype(str(font_path), 34) if font_path.exists() else ImageFont.load_default()
small_font = ImageFont.truetype(str(font_path), 24) if font_path.exists() else ImageFont.load_default()
draw.text((600, 35), "세 모드별 공격 성공률 총 480건", fill="#111111", font=title_font, anchor="ma")
left, top, right, bottom = 150, 120, 1100, 570
draw.line((left, top, left, bottom), fill="#555555", width=3)
draw.line((left, bottom, right, bottom), fill="#555555", width=3)
for pct in range(0, 101, 20):
    y = bottom - int((pct / 100) * (bottom - top))
    draw.line((left, y, right, y), fill="#D9D9D9", width=2)
    draw.text((left - 18, y), str(pct), fill="#444444", font=small_font, anchor="rm")
bar_width = 175
centers = [330, 625, 920]
colors = ["#8C2F39", "#B87922", "#24567A"]
for center, label, value, color in zip(centers, labels, values, colors):
    height = int((value / 100) * (bottom - top))
    draw.rectangle((center - bar_width//2, bottom - height, center + bar_width//2, bottom), fill=color)
    draw.text((center, max(top - 10, bottom - height - 18)), f"{value:.0f}%", fill="#111111", font=value_font, anchor="mb")
    draw.text((center, bottom + 28), label, fill="#111111", font=label_font, anchor="ma")
draw.text((38, 350), "공격 성공률 퍼센트", fill="#333333", font=small_font, anchor="mm")
image.save(chart_path)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    style = styles[style_name]
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style.font.color.rgb = RGBColor(0, 0, 0)
styles["Normal"].font.size = Pt(10.8)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.35
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.bold = True
styles["Heading 3"].font.size = Pt(11.5)
styles["Heading 3"].font.bold = True

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def add_table(headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, "24567A")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = font_name
            run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(font_size)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(cell)
    for r_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_index % 2:
                set_cell_shading(cells[i], "EEF4F8")
            for run in cells[i].paragraphs[0].runs:
                run.font.name = font_name
                run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
                run.font.size = Pt(font_size)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_margins(cells[i])
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_inline(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)

def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(9)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(95)
p.add_run("웹 애플리케이션의 인증·인가 우회 경로 분석 및 다중 보안 통제 적용에 따른 취약점 개선 효과 평가")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("16개 우회 경로의 취약 부분 강화 전체 강화 10회 반복 비교")
r.font.size = Pt(14)
r.bold = True
for text in ["작성자  김건희", "작성일  2026년 10월 12일", "로컬 실험 환경  Next.js PostgreSQL Prisma Playwright"]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18 if "작성자" in text else 5)
    p.add_run(text)
doc.add_page_break()

lines = source.splitlines()
skip_front = True
for raw in lines:
    line = raw.strip()
    if line == "## 초록":
        skip_front = False
    if skip_front or not line:
        continue
    if line.startswith("## "):
        heading = line[3:]
        doc.add_heading(heading, level=1)
        if heading == "4 결과":
            add_caption("표 6 모드별 10회 반복 결과")
            add_table(
                ["모드", "관찰 건수", "공격 성공", "공격 차단", "성공률", "중앙값 ms"],
                [
                    ["취약", 160, 160, 0, "100%", modes[0]["responseTimeMs"]["median"]],
                    ["인증·인가 부분 강화", 160, 80, 80, "50%", modes[1]["responseTimeMs"]["median"]],
                    ["전체 강화", 160, 0, 160, "0%", modes[2]["responseTimeMs"]["median"]],
                ],
                [0.9, 0.9, 0.9, 0.9, 0.8, 1.0]
            )
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(chart_path), width=Inches(5.7))
            add_caption("그림 1 모드별 공격 성공률")
            add_caption("표 7 범주별 공격 성공 결과")
            add_table(
                ["범주", "취약", "부분 강화", "전체 강화"],
                [
                    ["인증", "40/40", "0/40", "0/40"],
                    ["세션", "40/40", "40/40", "0/40"],
                    ["인가", "40/40", "0/40", "0/40"],
                    ["토큰 요청", "40/40", "40/40", "0/40"],
                ],
                [1.3, 1.5, 1.5, 1.2]
            )
            add_caption("표 8 강화 모드의 차단 상태 분포")
            add_table(
                ["HTTP 상태", "건수", "주요 의미"],
                [[400, 10, "쿠키 보호 조건 위반"], [401, 60, "인증 세션 토큰 무효"], [403, 50, "소유권 역할 출처 거부"], [410, 10, "소비된 재설정 토큰"], [422, 10, "비밀번호 정책 위반"], [429, 20, "로그인 시도 제한"]],
                [1.1, 0.9, 3.8]
            )
        continue
    if line.startswith("### "):
        heading = line[4:]
        doc.add_heading(heading, level=2)
        if heading == "1.3 질문을 가설로 좁힌 과정":
            add_caption("표 1 최종 선택에서 제외한 후보")
            add_table(
                ["후보", "접은 이유"],
                [
                    ["Docker와 로컬 실행 재현성", "환경 차이가 커 보안 통제 효과와 직접 연결되지 않음"],
                    ["비밀번호 정책 단독 비교", "세션과 권한 우회를 포괄하지 못해 연구 범위가 지나치게 좁음"],
                    ["OWASP 문헌 검토", "직접 수집한 원자료가 없어 실험 과제 목적을 충족하지 못함"],
                    ["알려진 모든 웹 공격", "일주일 안에 같은 조건으로 통제할 수 없는 과도한 범위"],
                ],
                [2.0, 4.7],
                8.5
            )
        elif heading == "3.2 변수와 통제":
            add_caption("표 2 실험 설계")
            add_table(
                ["구분", "내용"],
                [
                    ["독립 변수", "LAB_MODE 취약 인증인가 부분 강화 전체 강화"],
                    ["종속 변수", "공격 성공 여부와 성공률"],
                    ["기록 값", "회차 시나리오 모드 상태 응답시간 근거 관찰시각"],
                    ["고정 조건", "코드베이스 DB 구조 공격 입력 로컬 호스트 자동화 도구"],
                    ["반복과 표본", "16개 경로 × 3개 모드 × 10회 = 480건"],
                ],
                [1.25, 5.45]
            )
        elif heading == "3.3 인증·인가 우회 경로":
            add_caption("표 3 인증·인가 우회 경로와 강화 통제")
            scenario_rows = [
                ["AUTH-01", "사용자 열거", "균일 오류"], ["AUTH-02", "무차별 대입", "속도 제한"],
                ["AUTH-03", "약한 비밀번호", "길이와 차단 목록"], ["AUTH-04", "자격 증명 스터핑", "출처 누적 제한"],
                ["SESS-01", "세션 고정", "로그인 후 ID 재발급"], ["SESS-02", "쿠키 보호 누락", "HttpOnly Secure SameSite"],
                ["SESS-03", "로그아웃 토큰 재사용", "서버 측 폐기"], ["SESS-04", "만료 토큰", "만료시간 검증"],
                ["AUTHZ-01", "타인 객체 조회", "소유권 검사"], ["AUTHZ-02", "관리자 기능 우회", "역할 검사"],
                ["AUTHZ-03", "타인 객체 수정", "소유권 조건 수정"], ["AUTHZ-04", "타인 객체 삭제", "소유권 조건 삭제"],
                ["TOKEN-01", "JWT alg 변조", "허용 알고리즘 고정"], ["TOKEN-02", "보안 변경 전 토큰", "세션 일괄 폐기"],
                ["FLOW-01", "재설정 토큰 재사용", "해시 만료 단일 사용"], ["FLOW-02", "CSRF", "Origin과 CSRF 토큰"],
            ]
            add_table(["ID", "공격", "강화 통제"], scenario_rows, [1.0, 2.2, 3.5], 8.3)
        elif heading == "3.4 다중 보안 통제 적용 모델":
            add_caption("표 4 다중 보안 통제의 계층과 역할")
            add_table(
                ["계층", "주요 통제", "우회 경로에 미치는 역할"],
                [
                    ["예방", "균일 오류 비밀번호 정책 소유권 역할 JWT Origin CSRF 검증", "취약한 입력과 권한 결정을 보호 자산 도달 전에 거부"],
                    ["억제", "로그인 속도 제한 토큰과 세션 만료", "반복 기회와 인증 수단의 유효 시간을 축소"],
                    ["회복", "세션 ID 재발급 세션 폐기 단일 사용 재설정 토큰", "노출되거나 이전에 발급된 인증 수단의 재사용을 차단"],
                    ["감사", "시나리오 모드 결과 상태 근거의 DB 기록", "경로별 성공과 차단을 추적하고 재현 자료를 제공"],
                ],
                [0.8, 2.8, 3.1],
                8.2
            )
        elif heading == "3.5 취약점별 원인과 해결 근거":
            add_caption("표 5 취약점별 원인 해결책 선정 이유와 반복 결과")
            vulnerability_rows = [
                ["AUTH-01", "계정별 오류 노출", "동일한 401 오류", "유효 계정 식별 신호 제거", "10→0→0"],
                ["AUTH-02", "무제한 로그인 시도", "계정 출처별 429", "온라인 추측 기회 제한", "10→0→0"],
                ["AUTH-03", "짧고 흔한 암호 허용", "15자와 차단 목록", "약한 입력을 등록 시점에 거부", "10→0→0"],
                ["AUTH-04", "계정 순환 스터핑", "출처 누적 제한", "계정별 제한 우회 차단", "10→0→0"],
                ["SESS-01", "로그인 뒤 ID 유지", "난수 세션 재발급", "권한 변경 시 기존 연결 단절", "10→10→0"],
                ["SESS-02", "쿠키 보호 속성 누락", "HttpOnly Secure SameSite", "스크립트 전송 출처 위험 감소", "10→10→0"],
                ["SESS-03", "로그아웃 토큰 유효", "DB 세션 폐기", "복사된 토큰도 서버에서 거부", "10→10→0"],
                ["SESS-04", "만료 검증 누락", "요청마다 만료 검사", "사용 시점에 수명 강제", "10→10→0"],
                ["AUTHZ-01", "조회 소유권 미검사", "소유자 관리자 검사", "변조 가능한 객체 ID 불신", "10→0→0"],
                ["AUTHZ-02", "관리자 API 역할 미검사", "라우트 내부 ADMIN 검사", "직접 HTTP 호출 차단", "10→0→0"],
                ["AUTHZ-03", "수정 소유권 미검사", "수정 전 소유권 검사", "타인 객체 변경 차단", "10→0→0"],
                ["AUTHZ-04", "삭제 소유권 미검사", "삭제 전 소유권 검사", "타인 객체 삭제 차단", "10→0→0"],
                ["TOKEN-01", "공격자 alg 수락", "HS256 서명 iss aud exp", "서버 정책으로 토큰 검증", "10→10→0"],
                ["TOKEN-02", "암호 변경 전 세션 유효", "활성 세션 일괄 폐기", "탈취 토큰의 유효성 종료", "10→10→0"],
                ["FLOW-01", "재설정 토큰 재사용", "해시 만료 단일 사용", "노출 재전송 시간 초과 차단", "10→10→0"],
                ["FLOW-02", "자동 쿠키만으로 변경", "Origin과 CSRF 토큰", "요청 출처와 사용자 의도 확인", "10→10→0"],
            ]
            add_table(["ID", "취약 원인", "해결책", "선정 이유", "성공"], vulnerability_rows, [0.85, 1.45, 1.55, 2.3, 0.65], 7.5)
        continue
    if line.startswith("# "):
        continue
    p = doc.add_paragraph()
    add_inline(p, line.replace("  ", " "))

doc.add_heading("부록 A 시나리오별 반복 결과", level=1)
add_table(
    ["ID", "분류", "취약", "부분 강화", "전체 강화", "반복"],
    [[row["scenarioId"], row["category"], f'{row["vulnerableSuccesses"]}/10', f'{row["authAuthzOnlySuccesses"]}/10', f'{row["secureSuccesses"]}/10', row["repetitions"]] for row in summary["byScenario"]],
    [1.0, 1.3, 0.9, 1.1, 1.1, 0.7],
    8.4
)

doc.add_heading("부록 B T10 기준 대응표", level=1)
add_table(
    ["기준", "제출 위치", "상태"],
    [
        ["C01–C10", "1장 가설과 후보 선정 기록", "충족"],
        ["C11–C20", "2장 문헌 확인과 참고문헌", "충족"],
        ["C21–C30", "3장 원자료 실행 스크립트 재현 안내", "충족"],
        ["C31–C40", "4장 표 그래프 해석 한계", "충족"],
        ["C41–C51", "표지 초록 본문 참고문헌 ZIP AI 판단", "충족"],
    ],
    [1.2, 4.7, 0.8]
)

# Header and footer page number.
for sec in doc.sections:
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("웹 인증 보안 통제 비교 실험")
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

properties = doc.core_properties
properties.title = "웹 애플리케이션의 인증·인가 우회 경로 분석 및 다중 보안 통제 적용에 따른 취약점 개선 효과 평가"
properties.author = "김건희"
properties.subject = "16개 인증·인가 우회 경로와 다중 보안 통제의 취약점 개선 효과 비교"
properties.keywords = "인증 세션 권한 JWT CSRF Next.js PostgreSQL Playwright"

target = OUT / "auth-security-lab-paper.docx"
doc.save(target)
print(target)
